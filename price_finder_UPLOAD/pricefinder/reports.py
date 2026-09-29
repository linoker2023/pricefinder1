"""Вывод результатов: консольная таблица, CSV, XLSX, JSON, HTML-отчёт."""

from __future__ import annotations

import csv
import html as html_lib
import json
from datetime import datetime
from pathlib import Path
from typing import Any, Iterable, Sequence

from .parsers import Item, summarize
from .price import CURRENCY_CODES, format_amount

# ANSI-цвета (отключаются при --no-color или если вывод не в терминал)
RESET = "\033[0m"
BOLD = "\033[1m"
DIM = "\033[2m"
GREEN = "\033[32m"
YELLOW = "\033[33m"
CYAN = "\033[36m"
RED = "\033[31m"


def _c(text: str, color: str, enabled: bool) -> str:
    return f"{color}{text}{RESET}" if enabled else str(text)


def truncate(text: str, width: int) -> str:
    text = (text or "").replace("\n", " ").strip()
    if len(text) <= width:
        return text
    return text[: max(1, width - 1)] + "…"


# =========================================================================
#  Консоль
# =========================================================================

def print_table(
    items: Sequence[Item],
    *,
    query: str = "",
    limit: int = 20,
    color: bool = True,
    width: int = 120,
    base_currency: str = "RUB",
    show_url: bool = True,
) -> None:
    """Печатает топ предложений, отсортированный по цене (снизу вверх)."""
    if not items:
        print(_c("  Ничего не найдено.", YELLOW, color))
        return

    top = list(items[:limit])
    stats = summarize(top)
    sym = CURRENCY_CODES.get(base_currency, base_currency)
    lowest = top[0].normalized_price or top[0].price or 0
    median = stats.get("median") or 0

    if query:
        print(_c(f"\n  Самые низкие цены: «{query}»", BOLD, color))
    else:
        print(_c("\n  Самые низкие цены", BOLD, color))
    print(_c(f"  Предложений: {stats['count']} · минимум {fmt(lowest, sym)} · медиана {fmt(median, sym)}"
             f" · разброс {stats.get('spread_pct', 0)}%", DIM, color))
    print()

    w_rank, w_price, w_gain, w_site = 4, 15, 10, 20
    w_url = 34 if show_url else 0
    title_w = max(26, width - (w_rank + w_price + w_gain + w_site + w_url) - 3 * (4 if show_url else 3))

    head = (
        f"{'№':<{w_rank}}"
        f"{'Цена':>{w_price}}"
        f"{'Выгода':>{w_gain + 3}}"
        f"   {'Товар':<{title_w}}"
        f"   {'Площадка':<{w_site}}"
    )
    if show_url:
        head += f"   {'Ссылка':<{w_url}}"
    print(_c(head.rstrip(), BOLD, color))
    print(_c("─" * max(40, len(head)), DIM, color))

    for idx, it in enumerate(top, 1):
        price = it.normalized_price if it.normalized_price is not None else it.price
        price_s = fmt(price, sym)
        gain_s = ""
        if median and price:
            diff = (median - price) / median * 100
            if diff >= 0.5:
                gain_s = f"−{diff:.0f}%"
            elif diff <= -0.5:
                gain_s = f"+{abs(diff):.0f}%"

        title = truncate(it.title, title_w)
        if _is_out_of_stock(it.availability):
            title = _c(title, DIM, color) + _c(" · нет в наличии", RED, color)

        cells = [
            (f"{idx}.", w_rank, "left", DIM),
            (price_s, w_price, "right", GREEN if idx <= 3 else ""),
            (gain_s, w_gain, "right", GREEN if gain_s.startswith("−") else (RED if gain_s else "")),
            (title, title_w, "left", ""),
            (truncate(_site_label(it), w_site), w_site, "left", ""),
        ]
        if show_url:
            cells.append((_c(truncate(short_url(it.url), w_url), CYAN, color), w_url, "left", ""))

        line: list[str] = []
        for i, (text, w, align, col) in enumerate(cells):
            plain = _strip_ansi(text)
            if len(plain) > w:                       # на случай CJK/эмодзи
                text = truncate(plain, w)
                plain = _strip_ansi(text)
            gap = " " * max(0, w - len(plain))
            body = _c(text, col, color) if col and not text.startswith("\033") else text
            line.append((gap + body) if align == "right" else (body + gap))
        print(("   " if show_url else "  ").join(line).rstrip())
    print()


def _site_label(item: Item) -> str:
    if item.also_at:
        return f"{item.site_name} +{len(item.also_at)}"
    return item.site_name


def _is_out_of_stock(availability: str) -> bool:
    text = (availability or "").lower()
    return bool(text) and ("нет" in text or "out" in text or "отсут" in text)


def pad(text: str, width: int, align: str = "left", color: bool = True) -> str:
    text = str(text)
    gap = max(0, width - len(_strip_ansi(text)))
    return (" " * gap + text) if align == "right" else (text + " " * gap)


def _strip_ansi(text: str) -> str:
    import re

    return re.sub(r"\033\[[0-9;]*m", "", text or "")


def fmt(value: Any, symbol: str = "₽") -> str:
    if value is None:
        return "—"
    return f"{format_amount(value)} {symbol}".strip()


def short_url(url: str) -> str:
    if not url:
        return ""
    from urllib.parse import urlsplit

    parts = urlsplit(url)
    host = parts.netloc.replace("www.", "")
    path = parts.path.rstrip("/")
    return host + path if len(host + path) <= 60 else host + truncate(path, 40)


def print_site_summary(items: Sequence[Item], *, color: bool = True, base_currency: str = "RUB") -> None:
    stats = summarize(items)
    sym = CURRENCY_CODES.get(base_currency, base_currency)
    per_site = stats.get("sites") or {}
    if not per_site:
        return
    print(_c("  По площадкам:", BOLD, color))
    for name, st in sorted(per_site.items(), key=lambda kv: kv[1]["min"]):
        print(
            _c(f"    • {name:<18} предложений: {st['count']:<4} "
               f"от {fmt(st['min'], sym)} до {fmt(st['max'], sym)}", DIM, color)
        )
    print()


def print_errors(errors: Sequence[dict[str, str]], *, color: bool = True) -> None:
    if not errors:
        return
    print(_c("  Не удалось получить данные:", YELLOW, color))
    for err in errors:
        print(_c(f"    ! {err.get('site', '?')}: {err.get('error', '')[:120]}", DIM, color))
    print(_c("    Подсказка: добавьте js_render: true (нужен Playwright) или обновите селекторы в config/sites.yaml",
             DIM, color))
    print()


# =========================================================================
#  Файлы
# =========================================================================

def _rows(items: Sequence[Item]) -> list[dict[str, Any]]:
    rows = []
    for i, it in enumerate(items, 1):
        d = it.as_dict()
        d["rank"] = i
        rows.append(d)
    return rows


def save_csv(items: Sequence[Item], path: str | Path) -> Path:
    path = Path(path).expanduser()
    path.parent.mkdir(parents=True, exist_ok=True)
    rows = _rows(items)
    fields = ["rank", "site", "title", "price", "currency", "normalized_price", "price_text",
              "old_price", "availability", "rating", "reviews", "seller", "url", "strategy", "source_url"]
    with path.open("w", newline="", encoding="utf-8-sig") as fh:
        writer = csv.DictWriter(fh, fieldnames=fields, extrasaction="ignore", delimiter=";")
        writer.writeheader()
        writer.writerows(rows)
    return path


def save_json(items: Sequence[Item], path: str | Path, query: str = "") -> Path:
    path = Path(path).expanduser()
    path.parent.mkdir(parents=True, exist_ok=True)
    payload = {
        "query": query,
        "generated_at": datetime.now().isoformat(timespec="seconds"),
        "summary": summarize(items),
        "items": _rows(items),
    }
    path.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
    return path


def save_xlsx(items: Sequence[Item], path: str | Path, query: str = "") -> Path:
    path = Path(path).expanduser()
    path.parent.mkdir(parents=True, exist_ok=True)
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.utils import get_column_letter
    except ImportError as exc:
        raise RuntimeError("Для XLSX нужен openpyxl: pip install openpyxl (или используйте --csv/--json)") from exc

    wb = Workbook()
    ws = wb.active
    ws.title = "Цены"
    headers = ["#", "Площадка", "Товар", "Цена", "Валюта", "Цена (ед. валюта)", "Старая цена",
               "Наличие", "Рейтинг", "Отзывов", "Продавец", "Ссылка", "Источник", "Метод парсинга"]
    ws.append(headers)
    head_fill = PatternFill("solid", fgColor="1F4E79")
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = head_fill
        cell.alignment = Alignment(vertical="center")

    for row in _rows(items):
        ws.append([
            row["rank"], row["site"], row["title"], row["price"], row["currency"],
            row["normalized_price"], row.get("old_price"), row.get("availability"),
            row.get("rating"), row.get("reviews"), row.get("seller"),
            row["url"], row.get("source_url"), row.get("strategy"),
        ])
    for r in range(2, ws.max_row + 1):
        link = ws.cell(row=r, column=12).value
        if link:
            cell = ws.cell(row=r, column=12)
            cell.hyperlink = link
            cell.font = Font(color="0563C1", underline="single")
        ws.cell(row=r, column=4).number_format = "#,##0.00"
    best_fill = PatternFill("solid", fgColor="D8F0D8")
    for cell in ws[2]:
        cell.fill = best_fill
    widths = [4, 18, 60, 12, 8, 16, 12, 14, 9, 9, 20, 46, 40, 14]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

    stats = summarize(items)
    ws2 = wb.create_sheet("Сводка")
    ws2.append(["Запрос", query])
    ws2.append(["Сформировано", datetime.now().strftime("%Y-%m-%d %H:%M")])
    ws2.append(["Найдено предложений", stats.get("count", 0)])
    ws2.append(["Минимальная цена", stats.get("min")])
    ws2.append(["Медиана", stats.get("median")])
    ws2.append(["Средняя", stats.get("avg")])
    ws2.append(["Максимум", stats.get("max")])
    ws2.append(["Разброс, %", stats.get("spread_pct")])
    ws2.append([])
    ws2.append(["Площадка", "Предложений", "Мин.", "Макс."])
    for name, st in sorted((stats.get("sites") or {}).items(), key=lambda kv: kv[1]["min"]):
        ws2.append([name, st["count"], st["min"], st["max"]])
    for cell in ws2["A"]:
        cell.font = Font(bold=True)
    ws2.column_dimensions["A"].width = 26
    ws2.column_dimensions["B"].width = 22

    wb.save(path)
    return path


HTML_TEMPLATE = """<!DOCTYPE html>
<html lang="ru"><head><meta charset="utf-8">
<title>Цены: {query}</title>
<style>
  body{{font-family:-apple-system,Segoe UI,Roboto,Arial,sans-serif;margin:0;background:#f5f6f8;color:#1c1e21}}
  .wrap{{max-width:1180px;margin:0 auto;padding:28px 20px 60px}}
  h1{{font-size:24px;margin:0 0 4px}}
  .sub{{color:#6b7280;font-size:13px;margin-bottom:22px}}
  .cards{{display:flex;flex-wrap:wrap;gap:12px;margin-bottom:24px}}
  .card{{background:#fff;border:1px solid #e5e7eb;border-radius:12px;padding:14px 18px;min-width:150px}}
  .card b{{display:block;font-size:20px;margin-top:4px}}
  .card span{{font-size:12px;color:#6b7280}}
  table{{width:100%;border-collapse:collapse;background:#fff;border-radius:12px;overflow:hidden;
         box-shadow:0 1px 3px rgba(0,0,0,.06)}}
  th,td{{padding:10px 12px;text-align:left;font-size:13px;border-bottom:1px solid #eef0f3;vertical-align:top}}
  th{{background:#1f4e79;color:#fff;font-weight:600;position:sticky;top:0}}
  tr.best td{{background:#eaf7ea}}
  tr:hover td{{background:#f7fafc}}
  .price{{font-weight:700;white-space:nowrap;color:#0a7d32}}
  .gain{{color:#0a7d32;font-size:12px;white-space:nowrap}}
  .site{{white-space:nowrap;color:#374151}}
  a{{color:#1668c7;text-decoration:none}} a:hover{{text-decoration:underline}}
  .thumb{{width:46px;height:46px;object-fit:contain;border-radius:6px;background:#fff}}
  .title{{max-width:420px}}
  .muted{{color:#9aa1ab;font-size:11px}}
  .foot{{margin-top:18px;color:#9aa1ab;font-size:12px}}
</style></head><body><div class="wrap">
<h1>Самые низкие цены: «{query}»</h1>
<div class="sub">Сформировано {ts} · источников: {n_sites} · предложений: {n_items}</div>
<div class="cards">
  <div class="card"><span>Минимум</span><b>{min_price}</b></div>
  <div class="card"><span>Медиана</span><b>{median}</b></div>
  <div class="card"><span>Средняя</span><b>{avg}</b></div>
  <div class="card"><span>Максимум</span><b>{max_price}</b></div>
  <div class="card"><span>Разброс</span><b>{spread}%</b></div>
</div>
<table><thead><tr>
  <th>#</th><th>Цена</th><th>Выгода</th><th>Товар</th><th>Площадка</th><th>Наличие</th><th>Рейтинг</th><th>Ссылка</th>
</tr></thead><tbody>
{rows}
</tbody></table>
<div class="foot">pricefinder · данные собраны автоматически и могут отличаться от актуальных на сайте продавца. {errors}</div>
</div></body></html>
"""


def save_html(items: Sequence[Item], path: str | Path, query: str = "", errors: Sequence[dict] = (),
              base_currency: str = "RUB") -> Path:
    path = Path(path).expanduser()
    path.parent.mkdir(parents=True, exist_ok=True)
    stats = summarize(items)
    sym = CURRENCY_CODES.get(base_currency, base_currency)
    median = stats.get("median") or 0
    rows: list[str] = []
    for i, it in enumerate(items, 1):
        price = it.normalized_price if it.normalized_price is not None else it.price
        gain = ""
        if median and price:
            diff = (median - price) / median * 100
            if abs(diff) >= 0.5:
                gain = f"{'-' if diff > 0 else '+'}{abs(diff):.0f}% к медиане"
        img = f'<img class="thumb" src="{html_lib.escape(it.image)}" alt="" loading="lazy">' if it.image else ""
        rating = f"{it.rating:.1f}" + (f" <span class='muted'>({it.reviews})</span>" if it.reviews else "") if it.rating else "—"
        link = (f'<a href="{html_lib.escape(it.url)}" target="_blank" rel="noopener">открыть</a>'
                if it.url else "")
        rows.append(
            f"<tr class=\"{'best' if i <= 3 else ''}\">"
            f"<td>{i}</td>"
            f"<td class='price'>{fmt(price, sym)}</td>"
            f"<td class='gain'>{gain}</td>"
            f"<td class='title'>{html_lib.escape(it.title)}<div class='muted'>{html_lib.escape(it.seller)}</div></td>"
            f"<td class='site'>{img} {html_lib.escape(_site_label(it))}</td>"
            f"<td>{html_lib.escape(it.availability or '—')}</td>"
            f"<td>{rating}</td>"
            f"<td>{link}</td>"
            f"</tr>"
        )
    err_text = ""
    if errors:
        err_text = "Часть источников недоступна: " + ", ".join(
            f"{e.get('site')} ({e.get('error', '')[:60]})" for e in errors
        )
    html = HTML_TEMPLATE.format(
        query=html_lib.escape(query),
        ts=datetime.now().strftime("%d.%m.%Y %H:%M"),
        n_sites=len(stats.get("sites") or {}),
        n_items=stats.get("count", 0),
        min_price=fmt(stats.get("min"), sym),
        median=fmt(stats.get("median"), sym),
        avg=fmt(stats.get("avg"), sym),
        max_price=fmt(stats.get("max"), sym),
        spread=stats.get("spread_pct", 0),
        rows="\n".join(rows) or "<tr><td colspan='8'>Нет данных</td></tr>",
        errors=html_lib.escape(err_text),
    )
    path.write_text(html, encoding="utf-8")
    return path


def save_auto(items: Sequence[Item], path: str | Path, query: str = "", errors: Sequence[dict] = (),
              base_currency: str = "RUB") -> Path:
    """Сохраняет в формате по расширению: .csv / .json / .xlsx / .html"""
    path = Path(path).expanduser()
    suffix = path.suffix.lower()
    if suffix == ".csv":
        return save_csv(items, path)
    if suffix == ".json":
        return save_json(items, path, query)
    if suffix in (".xlsx", ".xlsm"):
        return save_xlsx(items, path, query)
    if suffix in (".html", ".htm"):
        return save_html(items, path, query, errors, base_currency)
    raise ValueError(f"Неподдерживаемый формат файла: {suffix} (csv/json/xlsx/html)")


# --- текстовый блок для уведомлений ----------------------------------------

def build_alert_html(events: Iterable[dict], query: str, base_currency: str = "RUB") -> tuple[str, str]:
    sym = CURRENCY_CODES.get(base_currency, base_currency)
    lines_html = ["<b>🔔 Снижение цены</b>", f"Запрос: <b>{html_lib.escape(query)}</b>", ""]
    lines_text = [f"🔔 Снижение цены по запросу «{query}»", ""]
    for ev in events:
        it: Item = ev["item"]
        price = it.normalized_price if it.normalized_price is not None else it.price
        prev = ev.get("prev_price")
        delta = f" (было {fmt(prev, sym)})" if prev else " (новая позиция)"
        kind = "минимум за всё время наблюдения" if ev.get("is_all_time_low") else "снижение"
        lines_html.append(
            f"• {fmt(price, sym)} — {kind}{delta}<br>"
            f"&nbsp;&nbsp;{html_lib.escape(truncate(it.title, 90))}<br>"
            f"&nbsp;&nbsp;<a href='{html_lib.escape(it.url)}'>{html_lib.escape(short_url(it.url))}</a> "
            f"[{html_lib.escape(it.site_name)}]"
        )
        lines_text.append(f"• {fmt(price, sym)}{delta} — {it.title[:90]} — {it.url} [{it.site_name}]")
    return "<br>".join(lines_html), "\n".join(lines_text)
